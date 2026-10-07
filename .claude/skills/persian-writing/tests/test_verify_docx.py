import os
import sys
import unittest

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if ROOT not in sys.path:
    sys.path.insert(0, ROOT)

try:  # python-docx is a dev-only dependency; the toolkit itself is stdlib-only
    import docx
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
except ImportError:  # pragma: no cover
    docx = None
from scripts.verify_docx import check


@unittest.skipIf(docx is None, "python-docx not installed")
class TestVerifyDocx(unittest.TestCase):
    def test_persianize_styles_inheritance(self):
        """Documents with style-level RTL and complex script definitions must inherit properly and pass verify_docx."""
        doc = docx.Document()
        # Set style-level RTL and complex script on Normal style
        st = doc.styles["Normal"]
        st.font.name = "Vazirmatn"
        rPr = st.element.get_or_add_rPr()
        rF = OxmlElement("w:rFonts")
        rF.set(qn("w:cs"), "Vazirmatn")
        rPr.append(rF)
        rPr.append(OxmlElement("w:rtl"))
        rPr.append(OxmlElement("w:szCs"))
        st.element.get_or_add_pPr().append(OxmlElement("w:bidi"))
        # Section bidi
        doc.sections[0]._sectPr.insert(0, OxmlElement("w:bidi"))

        doc.add_paragraph("متن تستی با ارثبری از استایل پیشفرض")
        path = "/tmp/test_persianize_styles_unit.docx"
        doc.save(path)

        errors, warnings, notes = check(path, expect_fonts=["Vazirmatn"])
        self.assertEqual(errors, [], f"Expected 0 errors with style inheritance, got: {errors}")
        self.assertTrue(
            any("styles.xml" in n for n in notes),
            f"Notes should indicate inheritance from styles.xml: {notes}",
        )

    def test_missing_rtl_and_cs(self):
        """Documents lacking RTL and complex-script font at both run and style levels must report hard errors."""
        doc = docx.Document()
        doc.add_paragraph("متن خام بدون هیچ تنظیمی")
        path = "/tmp/test_missing_rtl_unit.docx"
        doc.save(path)

        errors, warnings, notes = check(path, expect_fonts=[])
        self.assertTrue(any("no <w:rtl/>" in e for e in errors), "Must report missing <w:rtl/>.")
        self.assertTrue(any("no w:cs" in e for e in errors), "Must report missing w:cs font.")


if __name__ == "__main__":
    unittest.main()
