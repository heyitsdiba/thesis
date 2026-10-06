"""
Post-render mechanical lint for a rendered .pptx — the gate on the
composition-by-code path.

build-deck writes a per-deck pptxgenjs script, runs it, and then holds the
FILE (not the plan) to the brand identity. Everything the model could get
wrong between the spec and the slide surface is caught here: a run that
inherited its size from a master instead of naming one, a colour that never
resolved to a token, a font that is not the brand's, and — the reason this
module exists at all — text whose contrast is judged against the ground it
ACTUALLY sits on rather than against the paper token.

It reuses lint.py's rules rather than restating them, so the wording of a
geometry or palette violation stays coupled to the one mechanical gate:
check_colours, check_sizes, check_within_margins, check_no_overlap,
check_hierarchy, check_contrast and check_count all run against elements
reconstructed from the file. Only the file-shaped rules — explicit-ness,
ground resolution, background, bounds — are native to this module.

Public surface
--------------
composite(colour_hex, alpha, ground_hex) -> "#RRGGBB"
    Alpha-composite a colour over its ground. `alpha` is OOXML's
    <a:alpha val="N"/> (100000 = opaque); None means opaque.

lint_deck(pptx_path, brand, register=None) -> dict
    {"slides": [{"index", "violations", "notes", "elements"}, ...],
     "deck_notes": [...]}

main(argv=None) -> int
    CLI: lint_pptx.py <deck.pptx> --brand <brand.json> [--register R] [--json]
    0 clean, 1 violations, 2 the deck or the brand could not be read.

Shape walking
-------------
Shapes are visited in document order, which IS z-order (bottom to top) —
ground resolution depends on it. Group shapes are recursed into using each
CHILD's own box; pptxgenjs writes no groups, so that branch exists only so a
hand-edited or foreign deck does not silently skip half its shapes.

Duplicate suppression
---------------------
A single-run text shape yields the same [colour]/[size] message twice — once
from the run pass, once from the element pass, because the element's colour
and font_pt ARE that run's. Exactly-identical messages are collapsed per
slide: two identical lines tell the reader nothing the first did not.

Alpha
-----
python-pptx exposes no transparency, so run and fill alpha are read straight
off the XML: a:rPr/a:solidFill/a:srgbClr/a:alpha for a run,
spPr/a:solidFill/a:srgbClr/a:alpha for a shape fill.
"""

import argparse
import json
import os
import sys

from pptx import Presentation
from pptx.enum.dml import MSO_COLOR_TYPE, MSO_FILL
from pptx.enum.shapes import MSO_SHAPE_TYPE

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import lint  # noqa: E402
import tokens as tokens_mod  # noqa: E402

EMU_PER_INCH = 914400

# A mark (a rule, a dot, a connector, a bare panel) with no text near it is
# decoration nobody can read. "Near" is half an inch of rectangle gap in both
# axes; a label INSIDE the mark is a gap of zero, so it always counts.
LABEL_RADIUS_EMU = EMU_PER_INCH // 2

# Only the first N characters of a string identify it in a message.
SNIP = 30

_A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"

_FONT_HINT = "pass fontFace: T.fonts.heading or T.fonts.body"

_PICTURE_TYPES = (MSO_SHAPE_TYPE.PICTURE, MSO_SHAPE_TYPE.LINKED_PICTURE)


# ---------------------------------------------------------------------------
# Colour helpers
# ---------------------------------------------------------------------------

def composite(colour_hex, alpha, ground_hex):
    """Alpha-composite `colour_hex` over `ground_hex`, returning '#RRGGBB'.

    `alpha` is the OOXML <a:alpha val="N"/> value: 100000 is fully opaque,
    0 fully transparent, None (no alpha element) opaque. Per channel:
    round(a*c + (1-a)*g).
    """
    a = 1.0 if alpha is None else max(0.0, min(1.0, alpha / 100000.0))
    c = lint._hex_to_rgb(colour_hex)
    g = lint._hex_to_rgb(ground_hex)
    return "#" + "".join(
        f"{max(0, min(255, round(a * ci + (1.0 - a) * gi))):02X}"
        for ci, gi in zip(c, g)
    )


def _snip(text):
    """The first SNIP characters of a string — enough to identify it."""
    return str(text or "")[:SNIP]


def _rgb_hex(colour_format):
    """'#RRGGBB' when a ColorFormat carries an EXPLICIT RGB value, else None.

    A theme colour, an inherited colour, or no colour at all all return None:
    the brand gate wants the value named in the file, not one resolved from a
    master we did not write.
    """
    try:
        if colour_format is None or colour_format.type != MSO_COLOR_TYPE.RGB:
            return None
        return "#" + str(colour_format.rgb).upper()
    except Exception:  # noqa: BLE001 — an unreadable colour is simply absent
        return None


def _solid_fill_hex(fill):
    """'#RRGGBB' for an explicit solid RGB fill, else None."""
    try:
        if fill is None or fill.type != MSO_FILL.SOLID:
            return None
        return _rgb_hex(fill.fore_color)
    except Exception:  # noqa: BLE001
        return None


def _alpha_of(element):
    """The <a:alpha val="N"/> under an element's a:solidFill/a:srgbClr, or None."""
    if element is None:
        return None
    try:
        fill = element.find(_A + "solidFill")
        if fill is None:
            return None
        clr = fill.find(_A + "srgbClr")
        if clr is None:
            return None
        alpha = clr.find(_A + "alpha")
        if alpha is None:
            return None
        return int(alpha.get("val"))
    except Exception:  # noqa: BLE001
        return None


def _run_alpha(run):
    """A run's text transparency as an OOXML alpha value, or None."""
    try:
        return _alpha_of(run._r.find(_A + "rPr"))
    except Exception:  # noqa: BLE001
        return None


def _shape_fill_alpha(shape):
    """A shape fill's transparency as an OOXML alpha value, or None."""
    try:
        return _alpha_of(getattr(shape._element, "spPr", None))
    except Exception:  # noqa: BLE001
        return None


def _line_hex(shape):
    """'#RRGGBB' for an explicit solid line colour, else None."""
    try:
        line = shape.line
        if line.fill.type != MSO_FILL.SOLID:
            return None
        return _rgb_hex(line.color)
    except Exception:  # noqa: BLE001
        return None


def _background_hex(slide):
    """'#RRGGBB' when the slide carries an explicit solid background, else None."""
    try:
        return _solid_fill_hex(slide.background.fill)
    except Exception:  # noqa: BLE001
        return None


ABSENT, RGB, OTHER = "absent", "rgb", "other"
THEME = "theme"  # a run colour that is set but is not an RGB value


def _paint_state(fill):
    """(state, hex) for a FillFormat — paint is three-state.

    ABSENT: no fill object, no fill set, or an explicit no-fill. RGB: a solid
    fill with an explicit RGB colour (hex returned). OTHER: every other
    present representation — a solid theme/scheme colour, a gradient, a
    pattern, a picture, a texture. OTHER is visible paint the brand gate
    cannot read, so callers report it as [explicit]; only ABSENT is ignored
    (AR-006, AR-010).
    """
    if fill is None:
        return ABSENT, None
    try:
        kind = fill.type
    except Exception:  # noqa: BLE001
        return ABSENT, None
    if kind is None or kind == MSO_FILL.BACKGROUND:
        return ABSENT, None
    if kind == MSO_FILL.SOLID:
        hex_value = _rgb_hex(fill.fore_color)
        return (RGB, hex_value) if hex_value is not None else (OTHER, None)
    return OTHER, None


def _line_state(shape):
    """(state, hex) for a shape's line, by the same three-state rule."""
    try:
        line = shape.line
    except Exception:  # noqa: BLE001
        return ABSENT, None
    try:
        kind = line.fill.type
    except Exception:  # noqa: BLE001
        return ABSENT, None
    if kind is None or kind == MSO_FILL.BACKGROUND:
        return ABSENT, None
    if kind == MSO_FILL.SOLID:
        hex_value = _rgb_hex(line.color)
        return (RGB, hex_value) if hex_value is not None else (OTHER, None)
    return OTHER, None


def _colour_state(colour_format):
    """(state, hex) for a ColorFormat: ABSENT when unset, RGB, else OTHER."""
    try:
        if colour_format is None or colour_format.type is None:
            return ABSENT, None
        if colour_format.type == MSO_COLOR_TYPE.RGB:
            return RGB, "#" + str(colour_format.rgb).upper()
        return OTHER, None
    except Exception:  # noqa: BLE001
        return ABSENT, None


def _line_format_state(line_format):
    """(state, hex) for a LineFormat (a series or point outline), three-state."""
    try:
        kind = line_format.fill.type
    except Exception:  # noqa: BLE001
        return ABSENT, None
    if kind is None or kind == MSO_FILL.BACKGROUND:
        return ABSENT, None
    if kind == MSO_FILL.SOLID:
        hex_value = _rgb_hex(line_format.color)
        return (RGB, hex_value) if hex_value is not None else (OTHER, None)
    return OTHER, None


def _cell_fill_alpha(cell):
    """A table cell fill's <a:alpha> value, or None."""
    try:
        return _alpha_of(cell._tc.tcPr)
    except Exception:  # noqa: BLE001
        return None


# ---------------------------------------------------------------------------
# Geometry helpers
# ---------------------------------------------------------------------------

def _box_of(shape):
    """An (left, top, width, height) EMU dict for a shape; missing values as 0."""
    return {
        "left": int(shape.left or 0),
        "top": int(shape.top or 0),
        "width": int(shape.width or 0),
        "height": int(shape.height or 0),
    }


def _intersects(a, b):
    return (
        a["left"] < b["left"] + b["width"]
        and b["left"] < a["left"] + a["width"]
        and a["top"] < b["top"] + b["height"]
        and b["top"] < a["top"] + a["height"]
    )


def _within(inner, outer):
    return (
        inner["left"] >= outer["left"]
        and inner["top"] >= outer["top"]
        and inner["left"] + inner["width"] <= outer["left"] + outer["width"]
        and inner["top"] + inner["height"] <= outer["top"] + outer["height"]
    )


def _gap(a, b):
    """The rectangle gap between two boxes in EMU (0 when they touch/overlap)."""
    dx = max(0, max(a["left"], b["left"]) - min(
        a["left"] + a["width"], b["left"] + b["width"]))
    dy = max(0, max(a["top"], b["top"]) - min(
        a["top"] + a["height"], b["top"] + b["height"]))
    return dx, dy


# ---------------------------------------------------------------------------
# Shape walking
# ---------------------------------------------------------------------------

def _flatten(shapes, out):
    """Append every shape in z-order, recursing into groups by child box."""
    for shape in shapes:
        try:
            is_group = shape.shape_type == MSO_SHAPE_TYPE.GROUP
        except Exception:  # noqa: BLE001 — an unknown shape_type is not a group
            is_group = False
        if is_group:
            _flatten(shape.shapes, out)
        else:
            out.append(shape)
    return out


def _kind_of(shape):
    """Classify a shape into the element kind this lint reasons about."""
    if getattr(shape, "has_table", False):
        return "table"
    if getattr(shape, "has_chart", False):
        return "chart"
    try:
        if shape.shape_type in _PICTURE_TYPES:
            return "picture"
    except Exception:  # noqa: BLE001
        pass
    if getattr(shape, "has_text_frame", False):
        try:
            if shape.text_frame.text.strip():
                return "text"
        except Exception:  # noqa: BLE001
            pass
    box = _box_of(shape)
    if box["width"] <= 0 or box["height"] <= 0:
        return "connector"
    try:
        if shape.shape_type == MSO_SHAPE_TYPE.LINE:
            return "connector"
    except Exception:  # noqa: BLE001
        pass
    return "box"


# ---------------------------------------------------------------------------
# Run inspection
# ---------------------------------------------------------------------------

def _effective_size(run, paragraph):
    """A run's effective point size (its own, else its paragraph's), or None."""
    for font in (run.font, paragraph.font):
        try:
            if font.size is not None:
                return float(font.size.pt)
        except Exception:  # noqa: BLE001
            continue
    return None


def _effective_colour(run, paragraph):
    """A run's effective colour: '#RRGGBB', THEME, or None.

    The run's own setting wins whatever it is — an explicit theme colour on
    the run is THEME and never falls through to a paragraph's RGB (AR-010).
    Only an unset run inherits the paragraph's colour.
    """
    for font in (run.font, paragraph.font):
        state, hex_value = _colour_state(getattr(font, "color", None))
        if state == RGB:
            return hex_value
        if state == OTHER:
            return THEME
    return None


def _effective_font(run, paragraph):
    """A run's effective typeface name (its own, else its paragraph's), or None."""
    for font in (run.font, paragraph.font):
        try:
            if font.name:
                return font.name
        except Exception:  # noqa: BLE001
            continue
    return None


def _check_run(index, name, text, size, colour, alpha, font_name,
               ground, tk, brand_fonts):
    """Every rule a single run of text must clear. Returns violation messages."""
    messages = []
    label = f"[explicit] slide {index}: run '{_snip(text)}' in '{name}'"

    if size is None:
        messages.append(f"{label} has no explicit size; pass fontSize: T.scale.<step>")
    else:
        messages.extend(
            f"slide {index}: " + m for m in lint.check_sizes(
                [{"role": name, "text": text, "font_pt": size}], tk)
        )

    if colour is None:
        messages.append(f"{label} has no explicit colour; pass color: T.colours.<role>")
    elif colour == THEME:
        messages.append(
            f"{label} colour is a theme colour, not a token; pass color: T.colours.<role>"
        )
        colour = None
    else:
        messages.extend(
            f"slide {index}: " + m for m in lint.check_colours(
                [{"role": name, "text": text, "colour": colour}], tk)
        )

    if not font_name:
        messages.append(f"{label} has no explicit font; {_FONT_HINT}")
    elif font_name not in brand_fonts:
        messages.append(
            f"[font] slide {index}: run '{_snip(text)}' in '{name}' "
            f"uses font '{font_name}'; {_FONT_HINT}"
        )

    if size is not None and colour is not None and ground is not None:
        messages.extend(
            f"slide {index}: " + m for m in lint.check_contrast([{
                "role": name,
                "text": text,
                "colour": composite(colour, alpha, ground),
                "ground": ground,
                "font_pt": size,
            }], tk)
        )
    return messages


def _text_frame_runs(text_frame):
    """Yield (paragraph, run) for every run carrying non-blank text."""
    for paragraph in text_frame.paragraphs:
        for run in paragraph.runs:
            if run.text and run.text.strip():
                yield paragraph, run


# ---------------------------------------------------------------------------
# Chart inspection
# ---------------------------------------------------------------------------

def _check_chart_font(index, name, prop, size, colour, font_name, ground,
                      tk, brand_fonts):
    """The run rules, restated for a chart's Font object (which has no text)."""
    head = f"[chart] slide {index}: chart '{name}' {prop}"
    messages = []
    if size is None:
        messages.append(f"{head} [explicit] has no explicit size; "
                        f"pass fontSize: T.scale.<step>")
    elif size not in set(tk["type_scale"].values()):
        messages.append(f"{head} [size] has font_pt={size} which is not in type_scale")

    if colour is None:
        messages.append(f"{head} [explicit] has no explicit colour; "
                        f"pass color: T.colours.<role>")
    elif lint._norm(colour) not in _token_colours(tk):
        messages.append(f"{head} [colour] has colour={colour!r} "
                        f"which is not in colour_roles")

    if not font_name:
        messages.append(f"{head} [explicit] has no explicit font; {_FONT_HINT}")
    elif font_name not in brand_fonts:
        messages.append(f"{head} [font] uses font '{font_name}'; {_FONT_HINT}")

    if size is not None and colour is not None and ground is not None:
        messages.extend(f"{head} " + m for m in lint.check_contrast([{
            "role": name, "text": prop, "colour": colour,
            "ground": ground, "font_pt": size,
        }], tk))
    return messages


def _font_triple(font):
    """(size_pt, colour_hex, name) read off a chart Font, each None if absent."""
    size = None
    try:
        if font.size is not None:
            size = float(font.size.pt)
    except Exception:  # noqa: BLE001
        size = None
    return size, _rgb_hex(getattr(font, "color", None)), getattr(font, "name", None)


def _check_chart(index, shape, ground, tk, brand_fonts):
    """Hold a native chart's own text and series fills to the brand identity."""
    name = shape.name
    messages = []
    try:
        chart = shape.chart
    except Exception:  # noqa: BLE001 — an unreadable chart carries no text to gate
        return messages

    if getattr(chart, "has_title", False):
        try:
            frame = chart.chart_title.text_frame
        except Exception:  # noqa: BLE001
            frame = None
        if frame is not None:
            for paragraph, run in _text_frame_runs(frame):
                messages.extend(_check_chart_font(
                    index, name, f"title run '{_snip(run.text)}'",
                    _effective_size(run, paragraph),
                    _effective_colour(run, paragraph),
                    _effective_font(run, paragraph),
                    ground, tk, brand_fonts))

    for prop, getter in (("category axis tick labels", "category_axis"),
                         ("value axis tick labels", "value_axis")):
        try:
            axis = getattr(chart, getter)
            font = axis.tick_labels.font
        except Exception:  # noqa: BLE001 — this chart type has no such axis
            continue
        size, colour, font_name = _font_triple(font)
        messages.extend(_check_chart_font(index, name, prop, size, colour,
                                          font_name, ground, tk, brand_fonts))

    try:
        plots = list(chart.plots)
    except Exception:  # noqa: BLE001
        plots = []
    for i, plot in enumerate(plots):
        try:
            has_labels = plot.has_data_labels
        except Exception:  # noqa: BLE001
            has_labels = False
        if has_labels:
            size, colour, font_name = _font_triple(plot.data_labels.font)
            messages.extend(_check_chart_font(
                index, name, f"plot {i} data labels", size, colour,
                font_name, ground, tk, brand_fonts))
        try:
            series = list(plot.series)
        except Exception:  # noqa: BLE001
            series = []
        # AR-002/AR-005/AR-007/AR-008: the VISIBLE paint must be explicit and
        # a token — on the data points for pie/doughnut families, on the line
        # for line/radar/scatter families, on the fill for the rest — and any
        # other paint that is present must be a token too; theme or
        # non-solid paint anywhere is [explicit].
        family = getattr(chart.chart_type, "name", str(chart.chart_type)).upper()
        point_family = any(key in family for key in ("PIE", "DOUGHNUT"))
        line_family = any(key in family for key in ("LINE", "RADAR", "XY"))
        allowed = _token_colours(tk)

        def _judge(label, state, hex_value, required):
            if state == RGB and lint._norm(hex_value) not in allowed:
                return (f"[chart] slide {index}: chart '{name}' {label} [colour] "
                        f"has {hex_value!r} which is not in colour_roles")
            if state == OTHER:
                return (f"[chart] slide {index}: chart '{name}' {label} [explicit] "
                        f"is a theme or non-solid paint, not a token; pass a "
                        f"T.colours value")
            if state == ABSENT and required:
                return (f"[chart] slide {index}: chart '{name}' {label} [explicit] "
                        f"is automatic, not a token; pass chartColors with "
                        f"T.colours values")
            return None

        for j, one in enumerate(series):
            fill_state, fill_hex = _paint_state(one.format.fill)
            try:
                line_state, line_hex = _line_format_state(one.format.line)
            except Exception:  # noqa: BLE001
                line_state, line_hex = ABSENT, None

            if point_family:
                try:
                    points = list(one.points)
                except Exception:  # noqa: BLE001
                    points = []
                for k, point in enumerate(points):
                    p_state, p_hex = _paint_state(point.format.fill)
                    m = _judge(f"series {j} point {k} fill", p_state, p_hex, True)
                    if m:
                        messages.append(m)
                    # AR-011: a slice's own outline is visible paint too; an
                    # absent point line lets the series line govern.
                    try:
                        l_state, l_hex = _line_format_state(point.format.line)
                    except Exception:  # noqa: BLE001
                        l_state, l_hex = ABSENT, None
                    m = _judge(f"series {j} point {k} line", l_state, l_hex, False)
                    if m:
                        messages.append(m)
                required_fill, required_line = False, False
            else:
                required_fill, required_line = (not line_family), line_family

            for label, state, hex_value, required in (
                (f"series {j} fill", fill_state, fill_hex, required_fill),
                (f"series {j} line", line_state, line_hex, required_line),
            ):
                m = _judge(label, state, hex_value, required)
                if m:
                    messages.append(m)

    if getattr(chart, "has_legend", False):
        try:
            size, colour, font_name = _font_triple(chart.legend.font)
        except Exception:  # noqa: BLE001
            size = colour = font_name = None
        else:
            messages.extend(_check_chart_font(index, name, "legend", size, colour,
                                              font_name, ground, tk, brand_fonts))
    return messages


def _token_colours(tk):
    """The brand's token palette, normalised for membership tests."""
    return {lint._norm(v) for v in tk["colour_roles"].values()}


# ---------------------------------------------------------------------------
# Ground resolution
# ---------------------------------------------------------------------------

def _ground_beneath(box, panels, backdrop):
    """The colour a box sits on: the topmost earlier panel wholly holding it.

    Falls back to `backdrop` (the slide's own background, else the paper
    token) when nothing beneath contains the box.
    """
    for panel in reversed(panels):
        if _within(box, panel["box"]):
            return panel["effective"]
    return backdrop


def _straddled_panel(box, panels):
    """The topmost earlier panel this box intersects but does NOT sit inside."""
    touching = [p for p in panels if _intersects(box, p["box"])]
    if not touching:
        return None
    top = touching[-1]
    return None if _within(box, top["box"]) else top


# ---------------------------------------------------------------------------
# Per-slide pass
# ---------------------------------------------------------------------------

def _lint_slide(index, slide, tk, brand_fonts, slide_w, slide_h):
    """Return (violations, notes, elements) for one slide."""
    violations = []
    notes = []
    elements = []
    panels = []  # fill-only shapes seen so far, bottom-to-top

    paper = lint._norm(tk["colour_roles"].get("paper", "#FFFFFF"))
    background = _background_hex(slide)
    backdrop = lint._norm(background) if background else paper

    if background is not None and lint._norm(background) not in _token_colours(tk):
        violations.append(
            f"[background] slide {index}: background {lint._norm(background)} "
            f"is not a brand token"
        )
    else:
        try:
            bg_state = _paint_state(slide.background.fill)[0]
        except Exception:  # noqa: BLE001
            bg_state = ABSENT
        if bg_state == OTHER:
            violations.append(
                f"[background] slide {index}: background is a theme or non-solid "
                f"paint, not a token; pass background: {{ color: T.colours.<role> }}"
            )

    for shape in _flatten(slide.shapes, []):
        kind = _kind_of(shape)
        box = _box_of(shape)
        name = shape.name

        if kind in ("box", "connector"):
            # A python-pptx Connector (the fallback renderer's tree edges) has
            # no .fill at all; pptxgenjs lines are plain shapes and do.
            fill = getattr(shape, "fill", None)
            fill_state, fill_hex = _paint_state(fill)
            if fill_state == OTHER:
                violations.append(
                    f"[explicit] slide {index}: shape '{name}' fill is a theme or "
                    f"non-solid paint, not a token; pass fill: {{ color: T.colours.<role> }}"
                )
            if _line_state(shape)[0] == OTHER:
                violations.append(
                    f"[explicit] slide {index}: shape '{name}' line is a theme or "
                    f"non-solid paint, not a token; pass line: {{ color: T.colours.<role> }}"
                )
            element = {
                "kind": kind,
                "role": name,
                "text": "",
                "container": True,
                **box,
            }
            if fill_hex is not None:
                element["fill"] = fill_hex
            stroke = _line_hex(shape)
            if stroke is not None:
                element["stroke"] = stroke
            if (box["left"] <= 0 or box["top"] <= 0
                    or box["left"] + box["width"] >= slide_w
                    or box["top"] + box["height"] >= slide_h):
                element["full_bleed"] = True
            elements.append(element)
            if fill_hex is not None:
                beneath = _ground_beneath(box, panels, backdrop)
                panels.append({
                    "box": box,
                    "effective": composite(fill_hex, _shape_fill_alpha(shape), beneath),
                })
            continue

        if kind == "picture":
            elements.append({"kind": "picture", "role": name, "text": "", **box})
            continue

        # From here the shape carries text, so it needs a ground.
        straddled = _straddled_panel(box, panels)
        if straddled is not None and kind == "text":
            violations.append(
                f"[ground] slide {index}: text '{_snip(shape.text_frame.text)}' "
                f"straddles a panel edge; put it wholly on one ground"
            )
            ground = backdrop
        else:
            ground = _ground_beneath(box, panels, backdrop)

        if kind == "chart":
            elements.append({"kind": "chart", "role": name, "text": "", **box})
            violations.extend(_check_chart(index, shape, ground, tk, brand_fonts))
            continue

        if kind == "table":
            element = {"kind": "table", "role": name, "text": "", **box}
            fills = []
            table = shape.table
            for row in table.rows:
                for cell in row.cells:
                    cell_state, cell_fill = _paint_state(cell.fill)
                    if cell_fill is not None:
                        fills.append(cell_fill)
                    elif cell_state == OTHER:
                        violations.append(
                            f"[explicit] slide {index}: table '{name}' cell fill "
                            f"is a theme or non-solid paint, not a token; pass "
                            f"fill with a T.colours value"
                        )
                    # AR-009: a translucent cell fill sits on the table's ground.
                    cell_ground = (composite(cell_fill, _cell_fill_alpha(cell), ground)
                                   if cell_fill is not None else ground)
                    for paragraph, run in _text_frame_runs(cell.text_frame):
                        violations.extend(_check_run(
                            index, name, run.text,
                            _effective_size(run, paragraph),
                            _effective_colour(run, paragraph),
                            _run_alpha(run),
                            _effective_font(run, paragraph),
                            cell_ground, tk, brand_fonts))
            if fills:
                element["fills"] = fills
            elements.append(element)
            continue

        # kind == "text"
        frame = shape.text_frame
        # AR-001: a text box's own solid fill is what its runs sit on —
        # composited over the ground beneath when the fill carries alpha.
        own_state, own_fill = _paint_state(shape.fill)
        if own_fill is not None:
            ground = composite(own_fill, _shape_fill_alpha(shape), ground)
        elif own_state == OTHER:
            violations.append(
                f"[explicit] slide {index}: text box '{name}' fill is a theme or "
                f"non-solid paint, not a token; pass fill: {{ color: T.colours.<role> }}"
            )
        if _line_state(shape)[0] == OTHER:
            violations.append(
                f"[explicit] slide {index}: text box '{name}' line is a theme or "
                f"non-solid paint, not a token; pass line: {{ color: T.colours.<role> }}"
            )
        sizes = []
        first_colour = None
        for paragraph, run in _text_frame_runs(frame):
            size = _effective_size(run, paragraph)
            colour = _effective_colour(run, paragraph)
            if size is not None:
                sizes.append(size)
            if first_colour is None and colour is not None:
                first_colour = colour
            violations.extend(_check_run(
                index, name, run.text, size, colour, _run_alpha(run),
                _effective_font(run, paragraph), ground, tk, brand_fonts))

        element = {
            "kind": "text",
            "role": name,
            "text": frame.text,
            # A slides-ghost is decoration (the fallback renderer's ghost
            # numeral is the model): outside the hierarchy, like a tick or rule.
            "rank": (None if str(name).startswith("slides-ghost")
                     else 1 if (name == "slides-lead"
                                or str(name).startswith("slides-lead:"))
                     else 2),
            **box,
        }
        if sizes:
            element["font_pt"] = max(sizes)
        if first_colour is not None:
            element["colour"] = first_colour
        if own_fill is not None:
            element["fill"] = own_fill
        own_stroke = _line_hex(shape)
        if own_stroke is not None:
            element["stroke"] = own_stroke
        elements.append(element)

    # No declared lead: the largest text IS the lead, and says so.
    texts = [e for e in elements
             if e.get("kind") == "text" and e.get("font_pt") is not None
             and e.get("rank") is not None]  # ghosts stay outside the hierarchy
    if texts and not any(e.get("rank") == 1 for e in texts):
        lead = max(texts, key=lambda e: e["font_pt"])
        lead["rank"] = 1
        for other in texts:
            if other is not lead:
                other["rank"] = 2
        notes.append(
            f"slide {index}: no slides-lead declared; the largest text "
            f"('{_snip(lead['text'])}') was taken as the lead"
        )

    for check in (lint.check_colours, lint.check_sizes):
        violations.extend(f"slide {index}: " + m for m in check(elements, tk))
    violations.extend(
        f"slide {index}: " + m
        for m in lint.check_within_margins(elements, tk, slide_w, slide_h))
    for message in (lint.check_no_overlap(elements)
                    + lint.check_hierarchy(elements)
                    + lint.check_count(elements)):
        violations.append(f"slide {index}: " + message)

    slide_box = {"left": 0, "top": 0, "width": slide_w, "height": slide_h}
    for element in elements:
        if not _intersects(element, slide_box):
            violations.append(
                f"[bounds] slide {index}: element '{element['role']}' "
                f"lies wholly outside the slide"
            )

    # A mark nobody labelled is decoration.
    labels = [e for e in elements if e.get("kind") == "text"]
    for element in elements:
        if element.get("kind") not in ("box", "connector"):
            continue
        if any(max(_gap(element, label)) <= LABEL_RADIUS_EMU for label in labels):
            continue
        notes.append(
            f"slide {index}: orphan mark '{element['role']}' "
            f"has no label within 0.5 in"
        )

    seen = set()
    deduped = []
    for message in violations:
        if message not in seen:
            seen.add(message)
            deduped.append(message)
    return deduped, notes, elements


# ---------------------------------------------------------------------------
# Deck pass
# ---------------------------------------------------------------------------

DECK_NOTE_NO_GROUND = (
    "every slide sits on plain paper; the framing beats want a ground "
    "(see composing-in-code.md)"
)


def lint_deck(pptx_path, brand, register=None):
    """Lint a rendered .pptx against a brand profile.

    Returns {"slides": [{"index", "violations", "notes", "elements"}, ...],
    "deck_notes": [...]}. `elements` is the COUNT of elements reconstructed
    from that slide, not the elements themselves.
    """
    prs = Presentation(pptx_path)
    tk = tokens_mod.resolve_tokens(brand, None, register)
    brand_fonts = {v for v in (brand.get("fonts") or {}).values() if v}
    slide_w = int(prs.slide_width)
    slide_h = int(prs.slide_height)
    paper = lint._norm(tk["colour_roles"].get("paper", "#FFFFFF"))

    slides = []
    any_ground = False
    for i, slide in enumerate(prs.slides, start=1):
        violations, notes, elements = _lint_slide(
            i, slide, tk, brand_fonts, slide_w, slide_h)
        background = _background_hex(slide)
        if background is not None and lint._norm(background) != paper:
            any_ground = True
        # AR-003: a ground is a filled panel that bleeds; a bare line touching
        # the edge is not one.
        if any(e.get("kind") == "box" and e.get("full_bleed") and e.get("fill")
               for e in elements):
            any_ground = True
        slides.append({
            "index": i,
            "violations": violations,
            "notes": notes,
            "elements": len(elements),
        })

    deck_notes = [] if any_ground else [DECK_NOTE_NO_GROUND]
    return {"slides": slides, "deck_notes": deck_notes}


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------

def main(argv=None):
    """CLI entry point. 0 clean, 1 violations, 2 unreadable deck or brand."""
    parser = argparse.ArgumentParser(
        prog="lint_pptx.py",
        description="Hold a rendered .pptx to a brand's design tokens.",
    )
    parser.add_argument("deck", help="the rendered .pptx to lint")
    parser.add_argument("--brand", required=True, help="path to brand.json")
    parser.add_argument("--register", default=None,
                        help="the deck spec's register (keys the type scale)")
    parser.add_argument("--json", action="store_true", dest="as_json",
                        help="print the report as JSON instead of lines")
    args = parser.parse_args(argv)

    try:
        with open(args.brand, encoding="utf-8") as handle:
            brand = json.load(handle)
    except Exception as exc:  # noqa: BLE001 — any unreadable brand is exit 2
        print(f"error: {exc}", file=sys.stderr)
        return 2

    try:
        report = lint_deck(args.deck, brand, args.register)
    except Exception as exc:  # noqa: BLE001 — any unreadable deck is exit 2
        print(f"error: {exc}", file=sys.stderr)
        return 2

    violations = [v for s in report["slides"] for v in s["violations"]]
    dirty = [s for s in report["slides"] if s["violations"]]
    total_elements = sum(s["elements"] for s in report["slides"])

    if args.as_json:
        print(json.dumps(report, indent=2))
    else:
        for slide in report["slides"]:
            for message in slide["violations"]:
                print(message)
        for slide in report["slides"]:
            for note in slide["notes"]:
                print(note)
        for note in report["deck_notes"]:
            print(note)
        if violations:
            print(f"lint: {len(violations)} violation(s) on {len(dirty)} slide(s)")
        else:
            print(f"lint: clean ({len(report['slides'])} slides, "
                  f"{total_elements} elements)")

    return 1 if violations else 0


if __name__ == "__main__":
    sys.exit(main())
