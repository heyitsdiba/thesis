#!/usr/bin/env python3
"""Stamp a script-built deck so the revise round trip can find and read it.

    python3 stamp.py <deck>.pptx --spec <deck>.deck.md [--out <stamped>.pptx]

Writes the lineage comment render.py writes (`slides-spec: <basename>
sha256:<sha256 of the spec bytes>`) into core_properties.comments and names
each slide `slides-role:<role>` from the spec, in order — pptxgenjs cannot
name a slide, so this runs after the build script. Field shapes are named
`slides-field:<Field>` (or `slides-lead:<Field>` for the field that leads)
by the script itself, through pptxgenjs objectName.
"""
import argparse
import hashlib
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))


class StampError(Exception):
    """A deck and spec that cannot be stamped together."""


def lineage_comment(spec_path):
    """The lineage stamp string for a spec file — the one definition."""
    with open(spec_path, "rb") as fh:
        digest = hashlib.sha256(fh.read()).hexdigest()
    return f"slides-spec: {os.path.basename(spec_path)} sha256:{digest}"


def stamp_deck(pptx_path, spec_path, out_path=None):
    """Name every slide from the spec's roles and write the lineage comment.

    Returns the number of slides stamped. Raises StampError when the deck
    and the spec disagree on slide count.
    """
    from pptx import Presentation
    import render

    parsed = render.parse_spec(spec_path)
    prs = Presentation(pptx_path)
    if len(prs.slides) != len(parsed):
        raise StampError(
            f"deck has {len(prs.slides)} slide(s) but spec "
            f"{os.path.basename(spec_path)} has {len(parsed)}"
        )
    for slide, entry in zip(prs.slides, parsed):
        slide.name = f"slides-role:{entry['role']}"
    prs.core_properties.comments = lineage_comment(spec_path)
    prs.save(out_path or pptx_path)
    return len(parsed)


def main(argv=None):
    parser = argparse.ArgumentParser(
        description="Stamp a script-built deck for the revise round trip."
    )
    parser.add_argument("pptx", help="the built .pptx")
    parser.add_argument("--spec", required=True, help="the deck spec it was built from")
    parser.add_argument("--out", default=None, help="write here instead of in place")
    args = parser.parse_args(argv)
    try:
        count = stamp_deck(args.pptx, args.spec, args.out)
    except (OSError, StampError) as exc:
        print(f"error: {exc}", file=sys.stderr)
        return 1
    except Exception as exc:  # render.SpecError and python-pptx errors
        print(f"error: {exc}", file=sys.stderr)
        return 1
    print(f"stamped {count} slide(s) in {args.out or args.pptx}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
