"""Build a small test deck comparing the two approaches.

Slides 1-2 (option B): pre-rendered image; slide 2 zooms in using the Morph transition.
Slides 3-4 (option A): the cell as a native PowerPoint 3D model (.glb); slide 4 rotates/zooms
it, again with Morph. Older PowerPoint versions show the fallback image instead.

Usage: python3 make_test_deck.py RENDER.png MODEL.glb OUT.pptx
"""
import copy
import sys

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.opc.package import Part
from pptx.opc.packuri import PackURI
from pptx.util import Emu, Inches, Pt

RT_MODEL3D = "http://schemas.microsoft.com/office/2017/06/relationships/model3d"
NS = {
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
}

SLIDE_W, SLIDE_H = Inches(13.333), Inches(7.5)
TEXT = RGBColor(0x2B, 0x33, 0x4A)
MUTED = RGBColor(0x6B, 0x74, 0x8C)

MORPH = """<mc:AlternateContent xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006">
  <mc:Choice xmlns:p159="http://schemas.microsoft.com/office/powerpoint/2015/09/main" Requires="p159">
    <p:transition xmlns:p="{p}" xmlns:p14="http://schemas.microsoft.com/office/powerpoint/2010/main" spd="slow" p14:dur="2000">
      <p159:morph option="byObject"/>
    </p:transition>
  </mc:Choice>
  <mc:Fallback>
    <p:transition xmlns:p="{p}" spd="slow"><p:fade/></p:transition>
  </mc:Fallback>
</mc:AlternateContent>""".format(p=NS["p"])

# PowerPoint's default 3D-model lighting rig, camera looking down -z at the origin.
MODEL3D = """<mc:AlternateContent xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006">
 <mc:Choice xmlns:am3d="http://schemas.microsoft.com/office/drawing/2017/model3d" Requires="am3d">
  <p:graphicFrame xmlns:p="{p}" xmlns:a="{a}" xmlns:r="{r}">
   <p:nvGraphicFramePr>
    <p:cNvPr id="{id}" name="{name}" descr="3D model of a HEK293T cell (cutaway)"/>
    <p:cNvGraphicFramePr><a:graphicFrameLocks noGrp="1" noChangeAspect="1"/></p:cNvGraphicFramePr>
    <p:nvPr/>
   </p:nvGraphicFramePr>
   <p:xfrm><a:off x="{x}" y="{y}"/><a:ext cx="{cx}" cy="{cy}"/></p:xfrm>
   <a:graphic>
    <a:graphicData uri="http://schemas.microsoft.com/office/drawing/2017/model3d">
     <am3d:model3d r:embed="{rmodel}">
      <am3d:spPr>
       <a:xfrm><a:off x="{x}" y="{y}"/><a:ext cx="{cx}" cy="{cy}"/></a:xfrm>
       <a:prstGeom prst="rect"><a:avLst/></a:prstGeom>
      </am3d:spPr>
      <am3d:camera>
       <am3d:pos x="0" y="0" z="{camz}"/>
       <am3d:up dx="0" dy="36000000" dz="0"/>
       <am3d:lookAt x="0" y="0" z="0"/>
       <am3d:perspective fov="2700000"/>
      </am3d:camera>
      <am3d:trans>
       <am3d:meterPerModelUnit n="{mpu}" d="1000000"/>
       <am3d:preTrans dx="{pdx}" dy="{pdy}" dz="{pdz}"/>
       <am3d:scale><am3d:sx n="1000000" d="1000000"/><am3d:sy n="1000000" d="1000000"/><am3d:sz n="1000000" d="1000000"/></am3d:scale>
       <am3d:rot ax="{ax}" ay="{ay}" az="{az}"/>
       <am3d:postTrans dx="0" dy="0" dz="0"/>
      </am3d:trans>
      <am3d:raster rName="Office3DRenderer" rVer="16.0.8326"><am3d:blip r:embed="{rimg}"/></am3d:raster>
      <am3d:objViewport viewportSz="{vp}"/>
      <am3d:ambientLight><am3d:clr><a:scrgbClr r="50000" g="50000" b="50000"/></am3d:clr><am3d:illuminance n="500000" d="1000000"/></am3d:ambientLight>
      <am3d:ptLight rad="0"><am3d:clr><a:scrgbClr r="100000" g="75000" b="50000"/></am3d:clr><am3d:intensity n="9765625" d="1000000"/><am3d:pos x="21959998" y="70920001" z="16344003"/></am3d:ptLight>
      <am3d:ptLight rad="0"><am3d:clr><a:scrgbClr r="40000" g="60000" b="95000"/></am3d:clr><am3d:intensity n="12250000" d="1000000"/><am3d:pos x="-37964106" y="51130435" z="57631972"/></am3d:ptLight>
      <am3d:ptLight rad="0"><am3d:clr><a:scrgbClr r="86837" g="72700" b="100000"/></am3d:clr><am3d:intensity n="3125000" d="1000000"/><am3d:pos x="-37739005" y="0" z="-42532404"/></am3d:ptLight>
     </am3d:model3d>
    </a:graphicData>
   </a:graphic>
  </p:graphicFrame>
 </mc:Choice>
 <mc:Fallback>
  <p:pic xmlns:p="{p}" xmlns:a="{a}" xmlns:r="{r}">
   <p:nvPicPr>
    <p:cNvPr id="{id}" name="{name}" descr="3D model of a HEK293T cell (cutaway)"/>
    <p:cNvPicPr><a:picLocks noGrp="1" noRot="1" noChangeAspect="1"/></p:cNvPicPr>
    <p:nvPr/>
   </p:nvPicPr>
   <p:blipFill><a:blip r:embed="{rimg}"/><a:stretch><a:fillRect/></a:stretch></p:blipFill>
   <p:spPr><a:xfrm><a:off x="{x}" y="{y}"/><a:ext cx="{cx}" cy="{cy}"/></a:xfrm><a:prstGeom prst="rect"><a:avLst/></a:prstGeom></p:spPr>
  </p:pic>
 </mc:Fallback>
</mc:AlternateContent>"""


def add_text(slide, text, x, y, w, h, size, color=TEXT, bold=False):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = color
    p.font.name = "Helvetica Neue"
    return tb


def set_bg(slide, rgb=RGBColor(0xEE, 0xF1, 0xF6)):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = rgb


def add_morph(slide):
    sld = slide._element
    trans = etree.fromstring(MORPH)
    # transition goes right after clrMapOvr (before timing / extLst)
    clr = sld.find("p:clrMapOvr", NS)
    clr.addnext(trans)


def add_model3d(slide, glb_part, img_path, x, y, cx, cy, rot, cam_z, name):
    rmodel = slide.part.relate_to(glb_part, RT_MODEL3D)
    _, rimg = slide.part.get_or_add_image_part(img_path)
    sp_tree = slide.shapes._spTree
    next_id = max(int(e.get("id")) for e in sp_tree.iter("{%s}cNvPr" % NS["p"])) + 1 if list(
        sp_tree.iter("{%s}cNvPr" % NS["p"])) else 2
    xml = MODEL3D.format(
        p=NS["p"], a=NS["a"], r=NS["r"], id=next_id, name=name,
        x=int(x), y=int(y), cx=int(cx), cy=int(cy), rmodel=rmodel, rimg=rimg,
        camz=cam_z, mpu=MODEL["mpu"], pdx=MODEL["pre"][0], pdy=MODEL["pre"][1], pdz=MODEL["pre"][2],
        ax=rot[0], ay=rot[1], az=rot[2], vp=int(min(cx, cy)))
    sp_tree.insert(2, etree.fromstring(xml))  # behind the slide text (after nvGrpSpPr, grpSpPr)


# Model metadata (glTF is y-up): Blender bbox x[-10,9.7] y[-9,7.7] z[0,5.4]
# -> glTF centre (x, y, z) = (-0.15, 2.7, 0.65). preTrans is in millionths of a model unit.
MODEL = {"mpu": 50000, "pre": (150000, -2700000, -650000)}


def main(render, glb, out):
    prs = Presentation()
    prs.slide_width, prs.slide_height = SLIDE_W, SLIDE_H
    blank = prs.slide_layouts[6]

    # ---- Option B: pre-rendered image, Morph zoom -----------------------------------------
    s1 = prs.slides.add_slide(blank)
    set_bg(s1, RGBColor(0xEE, 0xF1, 0xF6))
    pic = s1.shapes.add_picture(render, 0, 0, SLIDE_W, SLIDE_H)
    pic.name = "!!CellRender"
    add_text(s1, "HEK293T — producer cell overview", Inches(0.5), Inches(0.35), Inches(8), Inches(0.6), 26, bold=True)
    add_text(s1, "Option B · pre-rendered (click to zoom with Morph)", Inches(0.5), Inches(0.95), Inches(8),
             Inches(0.4), 14, MUTED)

    s2 = prs.slides.add_slide(blank)
    set_bg(s2)
    z = 2.4  # zoom factor, centred on the nucleus (approx. 51% / 40% of the frame)
    fx, fy = 0.51, 0.40
    w, h = SLIDE_W * z, SLIDE_H * z
    pic = s2.shapes.add_picture(render, Emu(int(SLIDE_W / 2 - w * fx)), Emu(int(SLIDE_H / 2 - h * fy)), Emu(int(w)),
                                Emu(int(h)))
    pic.name = "!!CellRender"
    add_text(s2, "Nucleus — zoomed view", Inches(0.5), Inches(0.35), Inches(8), Inches(0.6), 26, bold=True)
    add_text(s2, "In the final deck this becomes a new close-up render or a short fly-in video", Inches(0.5),
             Inches(0.95), Inches(10), Inches(0.4), 14, MUTED)
    add_morph(s2)

    # ---- Option A: native 3D model ---------------------------------------------------------
    glb_part = Part(PackURI("/ppt/media/model3d1.glb"), "model/gltf-binary", prs.part.package,
                    open(glb, "rb").read())

    s3 = prs.slides.add_slide(blank)
    set_bg(s3)
    add_text(s3, "Option A · native 3D model", Inches(0.5), Inches(0.35), Inches(8), Inches(0.6), 26, bold=True)
    add_text(s3, "Click the cell and drag the centre handle to rotate it live", Inches(0.5), Inches(0.95),
             Inches(10), Inches(0.4), 14, MUTED)
    add_model3d(s3, glb_part, render, Inches(1.67), Inches(1.4), Inches(10), Inches(5.625),
                rot=(1800000, -2700000, 0), cam_z=67007152, name="!!Cell3D")

    s4 = prs.slides.add_slide(blank)
    set_bg(s4)
    add_text(s4, "Option A · Morph between two 3D views", Inches(0.5), Inches(0.35), Inches(9), Inches(0.6), 26,
             bold=True)
    add_text(s4, "Same model, rotated and enlarged — Morph animates the camera move", Inches(0.5),
             Inches(0.95), Inches(10), Inches(0.4), 14, MUTED)
    add_model3d(s4, glb_part, render, Inches(-3.3), Inches(-2.4), Inches(20), Inches(11.25),
                rot=(3600000, -1500000, 0), cam_z=67007152, name="!!Cell3D")
    add_morph(s4)

    prs.save(out)
    print("saved", out)


if __name__ == "__main__":
    main(*sys.argv[1:4])
